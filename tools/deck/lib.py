"""Shared layout + shape helpers for the AI Guardrails internal pitch deck.

Dark-theme tokens per solution-docs/40-sales-deck/VISUAL-DIRECTION.md §1 (all
placeholder values, pending a real brand kit). Every diagram built here is
native editable PowerPoint geometry with real text frames, per §9.
"""

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# --- tokens ---------------------------------------------------------------

BG = RGBColor(0x12, 0x16, 0x1B)
SURFACE = RGBColor(0x1D, 0x24, 0x2C)
SURFACE_ALT = RGBColor(0x16, 0x1C, 0x23)
HAIRLINE = RGBColor(0x2E, 0x39, 0x45)
TEXT = RGBColor(0xE8, 0xEC, 0xF2)
MUTED = RGBColor(0x98, 0xA4, 0xB3)
ACCENT = RGBColor(0x5A, 0xA0, 0xE0)
ACCENT_DIM = RGBColor(0x2B, 0x4A, 0x69)

# Semantic colours are bound to Decision domain actions only (VISUAL-DIRECTION §6).
ALLOW = RGBColor(0x4F, 0xB4, 0x77)
BLOCK = RGBColor(0xE0, 0x6C, 0x5F)
WARN = RGBColor(0xE0, 0xB3, 0x41)
LOG_ONLY = RGBColor(0x98, 0xA4, 0xB3)

HEAD_FONT = "Calibri"
BODY_FONT = "Calibri"

# --- grid -----------------------------------------------------------------

SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.75
CONTENT_W = SLIDE_W - 2 * MARGIN
TITLE_Y = 0.46
RULE_Y = 1.30
KICKER_Y = 1.42
BODY_Y = 1.95
BODY_BOTTOM = 6.62
FOOTER_Y = 6.80

# The standard two-column split of the body area. Slightly asymmetric because
# the left column almost always carries the longer text; both edges still land
# on MARGIN. Use these rather than restating the literals per slide.
SPLIT_L_X = MARGIN
SPLIT_L_W = 5.95
SPLIT_R_X = 7.13
SPLIT_R_W = 5.45


def _fmt(run, size=12, color=TEXT, bold=False, font=BODY_FONT, italic=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def textbox(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return box, tf


def para(tf, text, size=12, color=TEXT, bold=False, space_after=6, first=False,
         align=PP_ALIGN.LEFT, italic=False, font=BODY_FONT, line=1.15):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line
    run = p.add_run()
    run.text = text
    _fmt(run, size=size, color=color, bold=bold, font=font, italic=italic)
    return p


def slide_base(prs, title=None, kicker=None, footer=None):
    """Blank slide with the standard dark background, title block and footer."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG

    if title is not None:
        _, tf = textbox(s, MARGIN, TITLE_Y, CONTENT_W, 0.72)
        para(tf, title, size=30, bold=True, font=HEAD_FONT, first=True, space_after=0)
        rule = s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(MARGIN), Inches(RULE_Y), Inches(1.35), Inches(0.045)
        )
        rule.fill.solid()
        rule.fill.fore_color.rgb = ACCENT
        rule.line.fill.background()
        rule.shadow.inherit = False

    if kicker:
        _, tf = textbox(s, MARGIN, KICKER_Y, CONTENT_W, 0.40)
        para(tf, kicker, size=13.5, color=MUTED, first=True, space_after=0)

    if footer:
        _, tf = textbox(s, MARGIN, FOOTER_Y, CONTENT_W, 0.32)
        para(tf, footer, size=9, color=MUTED, first=True, space_after=0)

    return s


def notes(slide, sourcing, caveat, talk_track, question, prep=None):
    """Speaker notes carrying the four parts VISUAL-DIRECTION §8 requires.

    ``prep`` adds the optional fifth part §8 allows: what a presenter who did not
    write the deck has to establish or decide before this slide can be delivered
    to a team they do not know.
    """
    tf = slide.notes_slide.notes_text_frame
    tf.text = f"SOURCING — {sourcing}"
    parts = [
        ("CAVEAT NOT ON SLIDE", caveat),
        ("TALK TRACK", talk_track),
        ("IF ASKED", question),
    ]
    if prep is not None:
        parts.append(("PREP", prep))
    for label, body in parts:
        p = tf.add_paragraph()
        p.text = f"{label} — {body}"


def hide(slide):
    """Mark a slide as skipped in slideshow — presenter-only, never projected."""
    slide.element.set("show", "0")
    return slide


# --- shape primitives -----------------------------------------------------


def box(slide, x, y, w, h, text=None, fill=SURFACE, border=HAIRLINE, color=TEXT,
        size=11.5, bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE, border_w=1.0,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, pad=0.08):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if border is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = border
        sh.line.width = Pt(border_w)
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = 0.08
        except (IndexError, ValueError):
            pass
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    if text is not None:
        lines = text.split("\n")
        for i, line in enumerate(lines):
            bold_line = bold if i == 0 else False
            sz = size if i == 0 else size - 1.5
            col = color if i == 0 else MUTED
            para(tf, line, size=sz, color=col, bold=bold_line, first=(i == 0),
                 space_after=2, align=align)
    return sh


def _arrowhead(conn, end="tail"):
    ln = conn.line._get_or_add_ln()
    el = ln.makeelement(qn(f"a:{end}End"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(el)


def arrow(slide, x1, y1, x2, y2, color=ACCENT, width=1.25, dashed=False,
          connector=MSO_CONNECTOR.STRAIGHT):
    c = slide.shapes.add_connector(connector, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dashed:
        ln = c.line._get_or_add_ln()
        dash = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
        ln.append(dash)
    _arrowhead(c)
    return c


def edge_label(slide, x, y, w, text, size=9.5, color=MUTED, align=PP_ALIGN.CENTER):
    lines = text.split("\n")
    _, tf = textbox(slide, x, y, w, 0.20 + 0.18 * len(lines), align=align)
    for i, line in enumerate(lines):
        para(tf, line, size=size, color=color, first=(i == 0), space_after=0,
             align=align, line=1.1)


def chevron(slide, x, y, w, h, text, fill=SURFACE, border=HAIRLINE, color=TEXT,
            size=12, bold=True):
    return box(slide, x, y, w, h, text=text, fill=fill, border=border, color=color,
               size=size, bold=bold, shape=MSO_SHAPE.CHEVRON, pad=0.02)


def bullets(slide, x, y, w, items, size=13.5, gap=0.44, color=TEXT, dot=ACCENT,
            dot_size=0.085, sub_color=MUTED):
    """Custom bullet rows: a small accent square plus wrapped text."""
    cur = y
    for item in items:
        head, _, sub = item.partition("|")
        marker = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(cur + 0.085),
            Inches(dot_size), Inches(dot_size)
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = dot
        marker.line.fill.background()
        marker.shadow.inherit = False
        tw = w - 0.30
        _, tf = textbox(slide, x + 0.30, cur, tw, 0.4)
        para(tf, head.strip(), size=size, color=color, bold=bool(sub), first=True,
             space_after=1)
        if sub:
            para(tf, sub.strip(), size=size - 2, color=sub_color, space_after=0)
        cur += gap + (0.24 if sub else 0)
    return cur


def table(slide, x, y, w, h, data, col_w=None, header=True, size=11,
          header_size=11, first_col_bold=False, row_h=None):
    rows, cols = len(data), len(data[0])
    gt = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gt.table
    tbl.first_row = header
    tbl.horz_banding = False

    if col_w:
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Emu(int(Inches(w) * cw / total))
    if row_h:
        for i, rh in enumerate(row_h):
            tbl.rows[i].height = Inches(rh)

    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = SURFACE if (header and r == 0) else (
                SURFACE_ALT if r % 2 else BG)
            cell.margin_left = cell.margin_right = Inches(0.10)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            text, color, bold = val, TEXT, False
            if isinstance(val, tuple):
                text, color = val[0], val[1]
                bold = val[2] if len(val) > 2 else False
            if header and r == 0:
                color, bold = ACCENT, True
            elif first_col_bold and c == 0:
                bold = True
            para(tf, str(text), size=header_size if (header and r == 0) else size,
                 color=color, bold=bold, first=True, space_after=0, line=1.05)
    return tbl


_CHART_NS = ('xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
             'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"')


def _transparent(el):
    """Drop the default white fill and border from a chart element."""
    el.append(parse_xml(
        f'<c:spPr {_CHART_NS}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'))


def bar_chart(slide, x, y, w, h, categories, values, colors=None, cat_size=12,
              label_size=13, gap_width=70):
    """A native, editable horizontal bar chart on the deck's dark background.

    A real chart part rather than drawn rectangles, so the numbers stay bound to
    data a reviewer can open and check. The value axis is hidden — the bars are
    labelled with their own counts, and an axis of small integers reads as
    precision the underlying material does not have.
    """
    data = CategoryChartData()
    data.categories = categories
    data.add_series("count", values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y),
                                Inches(w), Inches(h), data)
    chart = gf.chart
    chart.has_title = False
    chart.has_legend = False

    plot = chart.plots[0]
    plot.gap_width = gap_width
    plot.vary_by_categories = False
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.font.size = Pt(label_size)
    labels.font.bold = True
    labels.font.color.rgb = TEXT
    labels.position = XL_LABEL_POSITION.OUTSIDE_END

    series = plot.series[0]
    for i, point in enumerate(series.points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = (
            colors[i] if colors else ACCENT)
        point.format.line.fill.background()

    cat_axis = chart.category_axis
    cat_axis.has_major_gridlines = False
    cat_axis.format.line.color.rgb = HAIRLINE
    cat_axis.tick_labels.font.size = Pt(cat_size)
    cat_axis.tick_labels.font.color.rgb = TEXT

    val_axis = chart.value_axis
    val_axis.visible = False
    val_axis.has_major_gridlines = False

    cs = chart._chartSpace
    cs.find(qn("c:chart")).addnext(parse_xml(
        f'<c:spPr {_CHART_NS}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'))
    _transparent(cs.find(qn("c:chart")).find(qn("c:plotArea")))
    return chart


def caption(slide, x, y, w, text, size=11.5, color=MUTED, align=PP_ALIGN.LEFT):
    _, tf = textbox(slide, x, y, w, 0.5, align=align)
    para(tf, text, size=size, color=color, first=True, space_after=0, align=align)


def panel(slide, x, y, w, h, fill=SURFACE_ALT, border=HAIRLINE):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = border
    sh.line.width = Pt(1.0)
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = 0.04
    except (IndexError, ValueError):
        pass
    return sh
