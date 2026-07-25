"""Shared slide-building helpers for the AI Guardrails proposal decks."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------- palette
INK = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x47, 0x55, 0x69)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
ACCENT2 = RGBColor(0x7C, 0x3A, 0xED)
GREEN = RGBColor(0x05, 0x96, 0x69)
AMBER = RGBColor(0xB4, 0x53, 0x09)
RED = RGBColor(0xB9, 0x1C, 0x1C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF1, 0xF5, 0xF9)
PANEL2 = RGBColor(0xE2, 0xE8, 0xF0)
CODEBG = RGBColor(0x0B, 0x12, 0x20)
CODEFG = RGBColor(0xE2, 0xE8, 0xF0)

FONT = "Segoe UI"
MONO = "Consolas"

SW = Inches(13.333)
SH = Inches(7.5)
MARGIN = Inches(0.62)
CONTENT_W = SW - 2 * MARGIN


def new_deck():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    return prs


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _box(slide, x, y, w, h, text="", size=14, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font=FONT, space_after=6, line=1.0,
         anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return tb


def _rect(slide, x, y, w, h, color, radius=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    _fill(shp, color)
    shp.shadow.inherit = False
    if radius is not None:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    shp.text_frame.text = ""
    return shp


def chrome(slide, prs, kicker, page=None, accent=ACCENT):
    """Thin top rule + running header + page number."""
    bar = _rect(slide, 0, 0, SW, Inches(0.075), accent)
    _box(slide, MARGIN, Inches(0.30), Inches(9.0), Inches(0.25),
         kicker.upper(), size=10, color=MUTED, bold=True)
    if page is not None:
        _box(slide, SW - MARGIN - Inches(1.2), Inches(0.30), Inches(1.2), Inches(0.25),
             str(page), size=10, color=MUTED, bold=True, align=PP_ALIGN.RIGHT)
    _box(slide, MARGIN, SH - Inches(0.45), Inches(9.0), Inches(0.25),
         "AI Guardrails — internal proposal  ·  derived from solution-docs/  ·  not for external distribution",
         size=8, color=MUTED)
    return bar


def title_slide(prs, kicker, title, subtitle, meta=""):
    s = _blank(prs)
    _rect(s, 0, 0, SW, SH, INK)
    _rect(s, 0, 0, Inches(0.14), SH, ACCENT)
    _box(s, Inches(1.1), Inches(1.55), Inches(11), Inches(0.4),
         kicker.upper(), size=13, color=RGBColor(0x93, 0xC5, 0xFD), bold=True)
    _box(s, Inches(1.1), Inches(2.05), Inches(11), Inches(1.9),
         title, size=44, color=WHITE, bold=True, line=1.05)
    _box(s, Inches(1.1), Inches(4.15), Inches(10), Inches(1.2),
         subtitle, size=17, color=RGBColor(0xCB, 0xD5, 0xE1), line=1.25)
    if meta:
        _box(s, Inches(1.1), SH - Inches(1.25), Inches(11), Inches(0.8),
             meta, size=11, color=MUTED, line=1.35)
    return s


def section_slide(prs, number, title, blurb, page=None):
    s = _blank(prs)
    _rect(s, 0, 0, SW, SH, INK)
    _rect(s, 0, 0, Inches(0.14), SH, ACCENT2)
    _box(s, Inches(1.1), Inches(2.35), Inches(2.0), Inches(1.4),
         number, size=64, color=ACCENT2, bold=True)
    _box(s, Inches(2.6), Inches(2.55), Inches(9.5), Inches(1.0),
         title, size=34, color=WHITE, bold=True)
    _box(s, Inches(2.65), Inches(3.55), Inches(9.0), Inches(1.2),
         blurb, size=15, color=RGBColor(0xCB, 0xD5, 0xE1), line=1.3)
    if page is not None:
        _box(s, SW - MARGIN - Inches(1.2), SH - Inches(0.6), Inches(1.2), Inches(0.25),
             str(page), size=10, color=MUTED, bold=True, align=PP_ALIGN.RIGHT)
    return s


def head(slide, title, sub=None):
    _box(slide, MARGIN, Inches(0.62), CONTENT_W, Inches(0.6),
         title, size=27, color=INK, bold=True)
    y = Inches(1.22)
    if sub:
        _box(slide, MARGIN, y, CONTENT_W, Inches(0.42), sub, size=13, color=SLATE, line=1.2)
        y = Inches(1.72)
    _rect(slide, MARGIN, y, Inches(1.05), Inches(0.035), ACCENT)
    return y + Inches(0.30)


def content_slide(prs, kicker, title, sub=None, page=None):
    s = _blank(prs)
    chrome(s, prs, kicker, page)
    y = head(s, title, sub)
    return s, y


def bullets(slide, x, y, w, items, size=14, gap=9, bullet_color=ACCENT):
    """items: list of (text, level) or (text, level, color)."""
    tb = slide.shapes.add_textbox(x, y, w, SH - y - Inches(0.7))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_top = 0
    first = True
    for it in items:
        text, level = it[0], it[1]
        col = it[2] if len(it) > 2 else (INK if level == 0 else SLATE)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap if level == 0 else gap - 3)
        p.line_spacing = 1.16
        if level == 0:
            r = p.add_run()
            r.text = "▪  "
            r.font.size = Pt(size)
            r.font.color.rgb = bullet_color
            r.font.name = FONT
            r.font.bold = True
        else:
            p.level = 1
            r = p.add_run()
            r.text = "      –  "
            r.font.size = Pt(size - 1)
            r.font.color.rgb = MUTED
            r.font.name = FONT
        r2 = p.add_run()
        r2.text = text
        r2.font.size = Pt(size if level == 0 else size - 1)
        r2.font.color.rgb = col
        r2.font.name = FONT
        r2.font.bold = (level == 0 and len(it) > 2)
    return tb


def table(slide, x, y, w, headers, rows, widths=None, size=11, head_size=11,
          head_bg=INK, row_h=Inches(0.34), head_h=Inches(0.38), zebra=True,
          col_colors=None):
    nrows, ncols = len(rows) + 1, len(headers)
    total_h = head_h + row_h * len(rows)
    gf = slide.shapes.add_table(nrows, ncols, x, y, w, total_h)
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    if widths:
        tot = sum(widths)
        for i, ww in enumerate(widths):
            tbl.columns[i].width = Emu(int(w * ww / tot))
    tbl.rows[0].height = head_h
    for i in range(1, nrows):
        tbl.rows[i].height = row_h

    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = head_bg
        cell.margin_left = Inches(0.09)
        cell.margin_right = Inches(0.06)
        cell.margin_top = Inches(0.04)
        cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = h
        r.font.size = Pt(head_size)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.name = FONT

    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL if (zebra and ri % 2 == 0) else WHITE
            cell.margin_left = Inches(0.09)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.line_spacing = 1.05
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(size)
            r.font.name = FONT
            colr = INK
            if col_colors and ci in col_colors:
                colr = col_colors[ci]
            r.font.color.rgb = colr
            r.font.bold = (ci == 0)
    return gf


def code_block(slide, x, y, w, h, code, size=11, label=None, fg=CODEFG, bg=CODEBG):
    _rect(slide, x, y, w, h, bg, radius=0.02)
    ty = y + Inches(0.16)
    if label:
        _box(slide, x + Inches(0.22), y + Inches(0.12), w - Inches(0.4), Inches(0.24),
             label, size=9, color=RGBColor(0x7D, 0xD3, 0xFC), bold=True, font=MONO)
        ty = y + Inches(0.44)
    tb = slide.shapes.add_textbox(x + Inches(0.22), ty, w - Inches(0.44), h - (ty - y) - Inches(0.14))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_top = 0
    lines = code.rstrip("\n").split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        p.line_spacing = 1.14
        r = p.add_run()
        r.text = ln if ln else " "
        r.font.size = Pt(size)
        r.font.name = MONO
        stripped = ln.strip()
        if stripped.startswith("#") or stripped.startswith("//"):
            r.font.color.rgb = RGBColor(0x86, 0xA0, 0xBC)
        elif stripped.startswith("$"):
            r.font.color.rgb = RGBColor(0x86, 0xEF, 0xAC)
        else:
            r.font.color.rgb = fg
    return tb


def cards(slide, y, items, h=Inches(2.5), gap=Inches(0.28), title_size=15, body_size=11.5):
    """items: list of dicts {tag, title, lines[], color}"""
    n = len(items)
    w = (CONTENT_W - gap * (n - 1)) / n
    for i, it in enumerate(items):
        x = MARGIN + (w + gap) * i
        col = it.get("color", ACCENT)
        _rect(slide, x, y, w, h, PANEL, radius=0.04)
        _rect(slide, x, y, w, Inches(0.05), col)
        _box(slide, x + Inches(0.22), y + Inches(0.24), w - Inches(0.44), Inches(0.24),
             it.get("tag", "").upper(), size=9.5, color=col, bold=True)
        _box(slide, x + Inches(0.22), y + Inches(0.54), w - Inches(0.44), Inches(0.5),
             it["title"], size=title_size, color=INK, bold=True, line=1.1)
        tb = slide.shapes.add_textbox(x + Inches(0.22), y + Inches(1.12),
                                      w - Inches(0.44), h - Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = 0
        tf.margin_top = 0
        for j, ln in enumerate(it["lines"]):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(5)
            p.line_spacing = 1.15
            r = p.add_run()
            r.text = "· " + ln
            r.font.size = Pt(body_size)
            r.font.color.rgb = SLATE
            r.font.name = FONT


def callout(slide, x, y, w, text, color=AMBER, size=11.5, label="CAVEAT", h=Inches(0.86)):
    _rect(slide, x, y, w, h, RGBColor(0xFE, 0xF3, 0xC7) if color == AMBER else PANEL, radius=0.05)
    _rect(slide, x, y, Inches(0.05), h, color)
    _box(slide, x + Inches(0.22), y + Inches(0.13), Inches(3.0), Inches(0.22),
         label, size=9, color=color, bold=True)
    _box(slide, x + Inches(0.22), y + Inches(0.38), w - Inches(0.5), h - Inches(0.5),
         text, size=size, color=INK, line=1.2)


def note(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def source(slide, text, y=None):
    _box(slide, MARGIN, y or (SH - Inches(0.78)), CONTENT_W, Inches(0.3),
         "Source: " + text, size=8.5, color=MUTED)


# public aliases (import * skips leading-underscore names)
rect = _rect
box = _box
blank = _blank
